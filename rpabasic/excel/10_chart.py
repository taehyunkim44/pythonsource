from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart

wb = load_workbook("./rpabasic/excel/range.xlsx")
ws = wb.active

# 막대 차트
# 차트 범위 설정
bar_value = Reference(ws, min_row=2, max_row=11, min_col=2,max_col=3)
# 차트 종류 설정
bar_chart = BarChart()
bar_chart.add_data(bar_value)

# 차트 삽입
ws.add_chart(bar_chart,"E1")


# LinceChart
line_value = Reference(ws, min_row=2, max_row=11, min_col=2,max_col=3)
# 차트 종류 설정
line_chart = LineChart()
line_chart.add_data(line_value)

# 차트 꾸미기
line_chart.title="성적표"
line_chart.style=20
line_chart.y_axis.title="점수"
line_chart.x_axis.title="번호"


# 차트 삽입
ws.add_chart(line_chart,"E15")

wb.save("./rpabasic/excel/range_chart.xlsx")
